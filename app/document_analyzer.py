"""
Analizador y gestor de carga programada de documentos municipales.

Funciones:
1. Analizar estructuralmente PDFs, planillas Excel, CSVs y Markdowns antes de indexar.
2. Encolar documentos para ingesta nocturna programada (03:00 AM) o inmediata.
3. Consultar y gestionar la cola de procesamiento del inbox.
"""
import os
import re
import json
import shutil
import datetime
from app import settings

COLA_JSON = os.path.join(settings.DATA_DIR, "inbox", ".cola_programada.json")


def _asegurar_directorios():
    inbox = os.path.join(settings.DATA_DIR, "inbox")
    os.makedirs(inbox, exist_ok=True)
    os.makedirs(settings.DATA_DIR, exist_ok=True)
    return inbox


def analizar_documento(file_path: str, filename: str) -> dict:
    """Analiza la estructura interna del documento y estima chunks y viabilidad."""
    ext = filename.lower().rsplit(".", 1)[-1]
    size_bytes = os.path.getsize(file_path) if os.path.exists(file_path) else 0
    size_kb = round(size_bytes / 1024, 1)

    resultado = {
        "valido": True,
        "nombre": filename,
        "extension": ext,
        "tamano_kb": size_kb,
        "paginas_o_filas": 0,
        "chunks_estimados": 1,
        "resumen": "",
        "detalles": [],
        "preview": ""
    }

    try:
        if ext == "pdf":
            import pymupdf
            doc = pymupdf.open(file_path)
            num_pages = doc.page_count
            total_chars = 0
            sample_text = ""

            for i in range(min(5, num_pages)):
                page_text = doc[i].get_text()
                total_chars += len(page_text)
                if not sample_text and page_text.strip():
                    sample_text = page_text.strip()[:400]

            # Si es un documento largo, proyectar el total de caracteres
            if num_pages > 5:
                avg_per_page = total_chars / 5
                total_chars = int(avg_per_page * num_pages)

            est_chunks = max(1, total_chars // settings.CHUNK_SIZE + 1)
            resultado["paginas_o_filas"] = num_pages
            resultado["chunks_estimados"] = est_chunks
            resultado["preview"] = sample_text

            if total_chars < 50:
                resultado["valido"] = False
                resultado["resumen"] = f"Aviso: El PDF tiene {num_pages} páginas pero no contiene texto seleccionable (parece ser una imagen escaneada sin OCR)."
            else:
                resultado["resumen"] = f"PDF analizado correctamente: {num_pages} páginas, ~{total_chars:,} caracteres. Se estiman ~{est_chunks} fragmentos (chunks) para el RAG."
            doc.close()

        elif ext in ("xlsx", "xls"):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheets = wb.sheetnames
            total_filas = 0
            detalles_hojas = []

            for sname in sheets:
                ws = wb[sname]
                rows = ws.max_row or 0
                cols = ws.max_column or 0
                total_filas += rows
                detalles_hojas.append(f"Hoja '{sname}': {rows} filas × {cols} cols")

            wb.close()
            # En Excel cada fila útil suele ser un chunk atómico o bloque
            est_chunks = max(1, total_filas // 2)
            resultado["paginas_o_filas"] = total_filas
            resultado["chunks_estimados"] = est_chunks
            resultado["detalles"] = detalles_hojas
            resultado["resumen"] = f"Planilla Excel analizada: {len(sheets)} hojas ('{', '.join(sheets[:3])}'), {total_filas} filas en total. Estimado de ~{est_chunks} registros."

        elif ext == "csv":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                lines = f.readlines()
            num_lines = len(lines)
            delimiter = ";" if ";" in (lines[0] if lines else "") else ","
            cols = len(lines[0].split(delimiter)) if lines else 0
            est_chunks = max(1, num_lines // 4)

            resultado["paginas_o_filas"] = num_lines
            resultado["chunks_estimados"] = est_chunks
            resultado["resumen"] = f"Archivo CSV estructurado: {num_lines} registros, {cols} columnas (separador '{delimiter}'). Estimado de ~{est_chunks} fragmentos."
            resultado["preview"] = "".join(lines[:4])

        elif ext == "md":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
            sections = len(re.findall(r"^\s{0,3}#{1,4}\s+", text, re.MULTILINE))
            est_chunks = max(1, sections or (len(text) // settings.CHUNK_SIZE + 1))
            resultado["paginas_o_filas"] = len(text.splitlines())
            resultado["chunks_estimados"] = est_chunks
            resultado["resumen"] = f"Documento Markdown limpio: {sections} secciones identificadas, ~{len(text):,} caracteres. Estimado de ~{est_chunks} fragmentos."
            resultado["preview"] = text[:350]

        else:
            resultado["valido"] = False
            resultado["resumen"] = f"Formato no soportado: .{ext}. Por favor sube archivos .pdf, .xlsx, .xls, .csv o .md."

    except Exception as e:
        resultado["valido"] = False
        resultado["resumen"] = f"Error al analizar el archivo: {e}"

    return resultado


def guardar_y_procesar(file_bytes: bytes, filename: str, programar: bool = True, categoria: str = "") -> dict:
    """Guarda el archivo en el inbox o lo indexa de inmediato."""
    inbox = _asegurar_directorios()
    clean_filename = re.sub(r"[^\w\.\-\_]", "_", filename)
    target_path = os.path.join(inbox, clean_filename)

    with open(target_path, "wb") as f:
        f.write(file_bytes)

    # Analizar estructura
    analisis = analizar_documento(target_path, clean_filename)
    if not analisis["valido"]:
        if os.path.exists(target_path):
            os.remove(target_path)
        return {"ok": False, "analisis": analisis, "mensaje": analisis["resumen"]}

    fecha_hora_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")

    if programar:
        # Registrar en la cola de procesamiento nocturno (03:00 AM)
        cola = obtener_cola_programada()
        # Evitar duplicados en la lista de cola
        cola = [item for item in cola if item.get("archivo") != clean_filename]
        cola.append({
            "archivo": clean_filename,
            "tamano": f"{analisis['tamano_kb']} KB",
            "categoria": categoria or "General",
            "chunks_estimados": analisis["chunks_estimados"],
            "hora_programada": "03:00 AM",
            "fecha_subida": fecha_hora_str,
            "resumen": analisis["resumen"]
        })
        with open(COLA_JSON, "w", encoding="utf-8") as f:
            json.dump(cola, f, ensure_ascii=False, indent=2)

        return {
            "ok": True,
            "modo": "programado",
            "analisis": analisis,
            "hora_programada": "03:00 AM",
            "mensaje": f"Documento '{clean_filename}' analizado exitosamente ({analisis['chunks_estimados']} chunks estimados). Ha quedado en cola para ser indexado automáticamente a las 03:00 AM."
        }
    else:
        # Indexar inmediatamente
        destino_final = os.path.join(settings.DATA_DIR, clean_filename)
        if os.path.abspath(target_path) != os.path.abspath(destino_final):
            shutil.move(target_path, destino_final)

        ext = clean_filename.lower().rsplit(".", 1)[-1]
        chunks_creados = 0
        from app import ingestion

        if ext == "pdf":
            chunks_creados = ingestion.ingest_pdf(clean_filename)
            ingestion.delete_original_file(clean_filename)
        elif ext in ("xlsx", "xls"):
            chunks_creados = ingestion.ingest_excel(clean_filename)
            ingestion.delete_original_file(clean_filename)
        elif ext == "md":
            chunks_creados = ingestion.ingest_markdown(clean_filename)
            ingestion.delete_original_file(clean_filename)
        elif ext == "csv":
            chunks_creados = ingestion.ingest_all()

        return {
            "ok": True,
            "modo": "inmediato",
            "analisis": analisis,
            "chunks_creados": chunks_creados,
            "mensaje": f"Documento '{clean_filename}' analizado e indexado de inmediato con éxito: {chunks_creados} fragmentos agregados a la base vectorial."
        }


def obtener_cola_programada() -> list:
    """Devuelve los documentos que están pendientes en data/inbox/."""
    inbox = _asegurar_directorios()
    if os.path.exists(COLA_JSON):
        try:
            with open(COLA_JSON, "r", encoding="utf-8") as f:
                data = json.load(f)
            # Filtrar solo los que realmente siguen en el disco
            return [d for d in data if os.path.exists(os.path.join(inbox, d.get("archivo", "")))]
        except Exception:
            pass

    # Si no hay JSON pero hay archivos en inbox
    archivos = [f for f in os.listdir(inbox) if not f.startswith(".")]
    cola = []
    for f in archivos:
        p = os.path.join(inbox, f)
        kb = round(os.path.getsize(p) / 1024, 1)
        cola.append({
            "archivo": f,
            "tamano": f"{kb} KB",
            "categoria": "General",
            "chunks_estimados": max(1, int(kb // 5)),
            "hora_programada": "03:00 AM",
            "fecha_subida": "Pendiente",
            "resumen": "Documento en cola para procesamiento"
        })
    return cola


def procesar_cola_inbox() -> dict:
    """Procesa todos los archivos en data/inbox/ ahora mismo."""
    inbox = _asegurar_directorios()
    archivos = [f for f in os.listdir(inbox) if not f.startswith(".")]
    if not archivos:
        return {"ok": True, "procesados": 0, "chunks": 0, "mensaje": "No hay documentos pendientes en la cola."}

    from app import ingestion
    total_chunks = 0
    procesados = 0

    for f in archivos:
        origen = os.path.join(inbox, f)
        destino = os.path.join(settings.DATA_DIR, f)
        shutil.move(origen, destino)

        ext = f.lower().rsplit(".", 1)[-1]
        try:
            if ext == "pdf":
                n = ingestion.ingest_pdf(f)
                ingestion.delete_original_file(f)
            elif ext in ("xlsx", "xls"):
                n = ingestion.ingest_excel(f)
                ingestion.delete_original_file(f)
            elif ext == "md":
                n = ingestion.ingest_markdown(f)
                ingestion.delete_original_file(f)
            elif ext == "csv":
                n = ingestion.ingest_all()
            else:
                n = 0
            total_chunks += n
            procesados += 1
        except Exception as e:
            print(f"[document_analyzer] Error al procesar {f}: {e}")

    # Limpiar JSON de cola
    if os.path.exists(COLA_JSON):
        os.remove(COLA_JSON)

    return {
        "ok": True,
        "procesados": procesados,
        "chunks": total_chunks,
        "mensaje": f"Se procesaron {procesados} documentos pendientes con un total de {total_chunks} fragmentos agregados a ChromaDB."
    }
