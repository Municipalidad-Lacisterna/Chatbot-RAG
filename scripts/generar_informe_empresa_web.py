#!/usr/bin/env python3
"""
Generador del documento formal:
Informe Técnico de Integración y Despliegue: Asistente Ciudadano Cisternin
Dirigido a la Empresa Externa de Desarrollo y Administración Web
Ilustre Municipalidad de La Cisterna
"""

import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="334155"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_callout(doc, text, title="NOTA TÉCNICA"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="32" w:space="0" w:color="1D4ED8"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"{title}: ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx(filename):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Header en páginas siguientes
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Municipalidad de La Cisterna · Informe de Integración Web")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Documento Técnico Institucional de Coordinación Externa — Uso Oficial")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Encabezado institucional
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_after = Pt(2)
    r1 = p_inst.add_run("ILUSTRE MUNICIPALIDAD DE LA CISTERNA\n")
    r1.bold = True
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    r2 = p_inst.add_run("DIRECCIÓN DE INFORMÁTICA Y TECNOLOGÍAS DE LA INFORMACIÓN\nPROYECTO SISTEMA DE ATENCIÓN CIUDADANA INTELIGENTE (CISTERNIN)")
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(18)
    p_line_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="0284C7"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(p_line_border)

    # Título principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(6)
    r_t = p_title.add_run("INFORME TÉCNICO DE INTEGRACIÓN Y DESPLIEGUE WEB: ASISTENTE CISTERNIN")
    r_t.bold = True
    r_t.font.name = "Arial"
    r_t.font.size = Pt(17)
    r_t.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(16)
    r_s = p_sub.add_run("Guía operacional de coordinación con la empresa proveedora de administración del portal municipal. Procedimientos de DNS, inserción de widget flotante y delimitación de responsabilidades de infraestructura.")
    r_s.font.name = "Arial"
    r_s.font.size = Pt(11)
    r_s.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Metadatos del documento
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Destinatario:", "Equipo de Desarrollo y Administración Web (Empresa Externa)"),
        ("Emisor Institucional:", "Área de Tecnologías e Infraestructura — Municipalidad de La Cisterna"),
        ("Fecha de Emisión:", "Octubre de 2026"),
        ("Materia Operativa:", "Integración de DNS, inserción de script y publicación de chatbot municipal")
    ]
    for idx, (label, val) in enumerate(meta_data):
        cell_l = meta_table.rows[idx].cells[0]
        cell_r = meta_table.rows[idx].cells[1]
        cell_l.width = Inches(2.2)
        cell_r.width = Inches(4.3)
        set_cell_background(cell_l, "F8FAFC")
        set_cell_background(cell_r, "F8FAFC")
        set_cell_margins(cell_l, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_r, top=60, bottom=60, left=100, right=100)
        
        pl = cell_l.paragraphs[0]
        pl.paragraph_format.space_after = Pt(2)
        rl = pl.add_run(label)
        rl.bold = True
        rl.font.name = "Arial"
        rl.font.size = Pt(9.5)
        rl.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        pr = cell_r.paragraphs[0]
        pr.paragraph_format.space_after = Pt(2)
        rr = pr.add_run(val)
        rr.font.name = "Arial"
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(13.5)
        r.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(11)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
        return h

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Arial"
            rb.font.size = Pt(10)
            rb.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Arial"
            rb.font.size = Pt(10)
            rb.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        return p

    # 1. OBJETIVO
    add_h1("1. Objetivo y Alcance Técnico")
    add_p("El propósito de este informe es formalizar los requerimientos operacionales dirigidos a la empresa responsable de la administración técnica del portal web de la Municipalidad de La Cisterna, con el fin de habilitar el acceso al Asistente Virtual Ciudadano Cisternin.")
    add_p("El sistema ha sido desarrollado bajo un modelo de microservicio 100% desacoplado. Por consiguiente, la empresa proveedora de la web no requiere instalar entornos de ejecución, modificar bases de datos, gestionar servidores ni asumir tareas de mantenimiento de backend. Su participación se limita estrictamente a la configuración de un registro DNS y a la inserción de una etiqueta de script en el frontend.")

    # 2. CUADRO DE ROLES
    add_h1("2. Matriz de Roles y Responsabilidades Técnicas")
    add_p("A continuación se presenta la delimitación formal de tareas entre el equipo técnico municipal y la empresa administradora del portal web:")

    table_roles = doc.add_table(rows=6, cols=3)
    table_roles.alignment = WD_TABLE_ALIGNMENT.CENTER
    role_headers = ["Componente Técnico", "Responsable", "Descripción Operacional"]
    for i, h in enumerate(role_headers):
        cell = table_roles.rows[0].cells[i]
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    roles_data = [
        ("Servidor y Backend", "Municipalidad de La Cisterna", "Despliegue de contenedores Docker, base de datos vectorial ChromaDB y lógica de IA en puerto 8000."),
        ("VPN y Acceso SSH", "Municipalidad de La Cisterna", "Configuración de red privada cifrada mediante Tailscale. La empresa web no administra llaves ni consolas del servidor."),
        ("Túnel de Red (Cloudflared)", "Municipalidad de La Cisterna", "Instalación de cloudflared en el servidor local y vinculación directa con el backend."),
        ("Registro DNS (CNAME)", "Empresa Administradora Web", "Creación de registro CNAME en el panel DNS de lacisterna.cl apuntando al túnel municipal."),
        ("Integración de Widget", "Empresa Administradora Web", "Inserción de script asíncrono en el pie de página de la web institucional.")
    ]
    for idx, (comp, resp, desc) in enumerate(roles_data, start=1):
        row = table_roles.rows[idx]
        for c_idx, val in enumerate([comp, resp, desc]):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F8FAFC" if idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if c_idx == 1 and "Municipalidad" in val:
                r.bold = True
                r.font.color.rgb = RGBColor(0x0D, 0x6E, 0x64)
            elif c_idx == 1:
                r.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3. REQUERIMIENTOS PARA LA EMPRESA WEB
    add_h1("3. Pasos Requeridos por parte de la Empresa Web")
    add_p("Para completar la puesta en marcha, la empresa proveedora del sitio web debe ejecutar exclusivamente las siguientes dos acciones (tiempo estimado: inferior a 5 minutos):")

    add_h2("Paso 1: Configuración de Registro DNS Institucional (1 minuto)")
    add_p("En el panel de administración DNS de lacisterna.cl (Cloudflare o proveedor de hosting institucional), registrar la siguiente entrada:")

    table_dns = doc.add_table(rows=5, cols=2)
    table_dns.alignment = WD_TABLE_ALIGNMENT.CENTER
    dns_rows = [
        ("Tipo de Registro:", "CNAME"),
        ("Nombre / Host:", "chatbot (para generar la URL pública chatbot.lacisterna.cl)"),
        ("Destino / Target:", "<ID_DEL_TUNEL>.cfargotunnel.com (proporcionado por la Municipalidad)"),
        ("Estado de Proxy:", "Proxied / Nube Naranja activada (Cifrado TLS automático)"),
        ("TTL:", "Automático")
    ]
    for idx, (label, val) in enumerate(dns_rows):
        cl = table_dns.rows[idx].cells[0]
        cr = table_dns.rows[idx].cells[1]
        cl.width = Inches(2.2)
        cr.width = Inches(4.3)
        set_cell_background(cl, "F8FAFC")
        set_cell_background(cr, "FFFFFF")
        set_cell_margins(cl, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cr, top=60, bottom=60, left=100, right=100)
        
        pl = cl.paragraphs[0]
        rl = pl.add_run(label)
        rl.bold = True
        rl.font.name = "Arial"
        rl.font.size = Pt(9.5)
        
        pr = cr.paragraphs[0]
        rr = pr.add_run(val)
        rr.font.name = "Arial"
        rr.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("Paso 2: Inserción del Widget en el Portal Institucional (2 minutos)")
    add_p("En la plantilla global o pie de página (footer) del portal web (React, Strapi o HTML), agregar la siguiente etiqueta:")
    add_code_block(doc, "<!-- Widget Oficial Cisternin - Municipalidad de La Cisterna -->\n<script \n  src=\"https://chatbot.lacisterna.cl/static/widget.js\" \n  data-bot-url=\"https://chatbot.lacisterna.cl/chat\" \n  defer>\n</script>")

    add_callout(doc, "Si por razones de política de diseño el portal web prefiere utilizar un botón estático institucional en lugar del widget flotante, pueden simplemente crear un botón en el menú de atención ciudadana con enlace directo hacia: https://chatbot.lacisterna.cl/chat", "ALTERNATIVA DE ENLACE DIRECTO")

    # 4. ESPECIFICACIONES TÉCNICAS DEL WIDGET
    add_h1("4. Especificaciones Técnicas y Aislamiento del Frontend")
    add_p("Para dar total tranquilidad al equipo de desarrollo web, el widget ha sido diseñado bajo los siguientes criterios de calidad:")
    add_bullet("El script tiene un tamaño inferior a 6 KB y se carga de manera diferida (defer), sin bloquear la renderización de la página ni afectar las métricas de rendimiento (Google PageSpeed / Core Web Vitals).", "Cero impacto en el rendimiento web: ")
    add_bullet("La interfaz de conversación se ejecuta dentro de un elemento Iframe aislado con Sandbox. Los estilos CSS y scripts de Cisternin no tienen acceso ni colisionan con las hojas de estilo del sitio municipal.", "Aislamiento visual absoluto (CSS Sandbox): ")
    add_bullet("El componente responde fluidamente a resoluciones móviles, tablets y monitores de escritorio.", "Compatibilidad multiplataforma: ")

    # 5. LABORES MUNICIPALES (CÓMO ENCAJA LA INFRAESTRUCTURA)
    add_h1("5. Infraestructura y Operaciones a cargo de la Municipalidad")
    add_p("A título informativo, se expone la arquitectura que implementa internamente el equipo de TI municipal para sostener el servicio:")
    
    add_bullet("El equipo municipal configura Tailscale directamente en el servidor. Esto crea una interfaz de red privada virtual con dirección IP en el segmento 100.x.y.z. Desde allí, el personal de informática y jefatura administra el servidor mediante SSH y accede a las consolas de gestión (/admin y /docs). La empresa web queda 100% desvinculada de la gestión de puertos SSH, usuarios de sistema o firewalls.", "Gestión de Red Privada y SSH (Tailscale): ")
    
    add_bullet("El equipo municipal ejecuta el servicio oficial cloudflared en el servidor de producción. Este servicio abre un túnel criptográfico saliente hacia la red de Cloudflare y asocia el puerto local 8000 con el subdominio municipal configurado en el CNAME. Gracias a este esquema, el municipio no abre puertos 80 ni 443 en su router institucional.", "Gestión del Túnel de Publicación (Cloudflare Tunnel): ")

    # 6. FLUJO DE VALIDACIÓN
    add_h1("6. Flujo de Coordinación y Salida a Producción")
    add_p("El procedimiento de activación conjunta se resume en los siguientes tres pasos:")
    add_p("1. La Municipalidad de La Cisterna proporciona a la empresa web el identificador de destino del túnel (ejemplo: cadcff80-6846-4eff-afc5-2b2bbf0c3d03.cfargotunnel.com).")
    add_p("2. La Empresa Web crea el registro CNAME en el DNS e inserta el script del widget en el portal institucional.")
    add_p("3. Validación conjunta en línea: Se comprueba la visualización de la burbuja interactiva de Cisternin en www.lacisterna.cl y la correcta recepción de consultas ciudadanas.")

    # Firma
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.keep_with_next = True
    r_s1 = p_sign.add_run("ÁREA DE DESARROLLO E INFRAESTRUCTURA TI\n")
    r_s1.bold = True
    r_s1.font.name = "Arial"
    r_s1.font.size = Pt(10)
    r_s1.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)
    
    r_s2 = p_sign.add_run("Proyecto Asistente Virtual Ciudadano Cisternin\nIlustre Municipalidad de La Cisterna — Octubre de 2026")
    r_s2.font.name = "Arial"
    r_s2.font.size = Pt(9.5)
    r_s2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.save(filename)
    print(f"Documento Word guardado exitosamente: {filename}")

if __name__ == "__main__":
    out_docx = "/home/aspen/Alcaldia_Practica/docs/Informe_Tecnico_Integracion_Empresa_Web.docx"
    build_docx(out_docx)
    
    static_docx = "/home/aspen/Alcaldia_Practica/static/docs/Informe_Tecnico_Integracion_Empresa_Web.docx"
    os.makedirs(os.path.dirname(static_docx), exist_ok=True)
    subprocess.run(["cp", out_docx, static_docx], check=True)
    
    # Copiar también a /home/aspen/
    subprocess.run(["cp", out_docx, "/home/aspen/Informe_Tecnico_Integracion_Empresa_Web.docx"], check=True)
    
    print("Convirtiendo a PDF con LibreOffice...")
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to", "pdf",
        out_docx,
        "--outdir", "/home/aspen/Alcaldia_Practica/docs"
    ]
    subprocess.run(cmd, check=True)
    
    out_pdf = "/home/aspen/Alcaldia_Practica/docs/Informe_Tecnico_Integracion_Empresa_Web.pdf"
    static_pdf = "/home/aspen/Alcaldia_Practica/static/docs/Informe_Tecnico_Integracion_Empresa_Web.pdf"
    subprocess.run(["cp", out_pdf, static_pdf], check=True)
    subprocess.run(["cp", out_pdf, "/home/aspen/Informe_Tecnico_Integracion_Empresa_Web.pdf"], check=True)
    print("PDF generado y copiado exitosamente a static/docs/ y /home/aspen/")
