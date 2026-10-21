#!/usr/bin/env python3
"""
Generador del documento formal:
Guía Técnica de Infraestructura de Red, Acceso Remoto y Despliegue en Producción
Chatbot Municipal Cisternin - Ilustre Municipalidad de La Cisterna
"""

import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Borde izquierdo azul grisáceo
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="334155"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_callout(doc, text, title="NOTA TÉCNICA"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="32" w:space="0" w:color="1D4ED8"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"{title}: ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx(filename):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Header en páginas siguientes
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Municipalidad de La Cisterna · Guía Técnica de Red y Despliegue")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Documento Técnico Institucional — Confidencialidad y Uso Interno")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Encabezado institucional
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_after = Pt(2)
    r1 = p_inst.add_run("ILUSTRE MUNICIPALIDAD DE LA CISTERNA\n")
    r1.bold = True
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    r2 = p_inst.add_run("DIRECCIÓN DE INFORMÁTICA Y TECNOLOGÍAS DE LA INFORMACIÓN\nPROYECTO SISTEMA DE ATENCIÓN CIUDADANA INTELIGENTE (CISTERNIN)")
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(20)
    p_line_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="0284C7"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(p_line_border)

    # Título principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(6)
    r_t = p_title.add_run("GUÍA TÉCNICA DE INFRAESTRUCTURA DE RED, ACCESO REMOTO Y DESPLIEGUE EN PRODUCCIÓN")
    r_t.bold = True
    r_t.font.name = "Arial"
    r_t.font.size = Pt(18)
    r_t.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    r_s = p_sub.add_run("Manual de Operación de Doble Red Cero-Confianza (Tailscale y Cloudflare Tunnel), Procedimientos de Conexión SSH para Jefatura, Despliegue con/sin Script y Coordinación con el Área Web.")
    r_s.font.name = "Arial"
    r_s.font.size = Pt(11)
    r_s.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Metadatos del documento
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Fecha de Elaboración:", "Octubre de 2026"),
        ("Versión del Documento:", "1.0 (Definitiva para Producción)"),
        ("Destinatarios:", "Administración Municipal, Dirección de Informática, Jefatura de Proyectos"),
        ("Alcance Operativo:", "Servidor Municipal de Producción, Conectividad de Gestión y Portal Web")
    ]
    for idx, (label, val) in enumerate(meta_data):
        cell_l = meta_table.rows[idx].cells[0]
        cell_r = meta_table.rows[idx].cells[1]
        cell_l.width = Inches(2.2)
        cell_r.width = Inches(4.3)
        set_cell_background(cell_l, "F8FAFC")
        set_cell_background(cell_r, "F8FAFC")
        set_cell_margins(cell_l, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_r, top=60, bottom=60, left=100, right=100)
        
        pl = cell_l.paragraphs[0]
        pl.paragraph_format.space_after = Pt(2)
        rl = pl.add_run(label)
        rl.bold = True
        rl.font.name = "Arial"
        rl.font.size = Pt(9.5)
        rl.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        pr = cell_r.paragraphs[0]
        pr.paragraph_format.space_after = Pt(2)
        rr = pr.add_run(val)
        rr.font.name = "Arial"
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Helper para encabezados
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(11.5)
        r.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
        return h

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Arial"
            rb.font.size = Pt(10)
            rb.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Arial"
            rb.font.size = Pt(10)
            rb.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        return p

    # -------------------------------------------------------------
    # 1. INTRODUCCIÓN Y PRINCIPIOS DE ARQUITECTURA
    # -------------------------------------------------------------
    add_h1("1. Introducción y Arquitectura de Red de Doble Capa")
    add_p("El presente documento técnico establece los procedimientos operacionales y estándares de conectividad requeridos para la puesta en marcha, gestión segura y publicación del sistema de asistencia ciudadana inteligente Cisternin en la infraestructura municipal.")
    add_p("Para dar cumplimiento a los estándares de seguridad de la información y soberanía operativa, el sistema implementa una arquitectura desacoplada de doble perímetro (Cero Confianza o Zero Trust), dividida en dos redes completamente estancas:")
    
    add_bullet(" Red de malla privada basada en protocolo criptográfico WireGuard. Su objetivo exclusivo es permitir que los administradores municipales y personal técnico de informática puedan acceder al servidor de producción, ejecutar tareas de mantenimiento por SSH y acceder a los paneles de control (/admin y /docs) de manera remota y segura, sin exponer puertos administrativos a internet abierto.", "1. Red Privada de Gestión (Tailscale Mesh VPN):")
    add_bullet(" Túnel cifrado de salida perimetral. Conecta el backend de atención ciudadana (puerto local 8000) con la infraestructura perimetral segura, asignando el subdominio oficial (chatbot.lacisterna.cl) y proveyendo terminación SSL/HTTPS y protección automática contra ataques de denegación de servicio (DDoS).", "2. Red Pública de Atención Ciudadana (Cloudflare Tunnel):")

    add_callout(doc, "Esta arquitectura no requiere abrir ningún puerto entrante (ni puerto 80, ni 443, ni 22) en el firewall o router perimetral de la municipalidad. Todas las conexiones se establecen mediante túneles seguros salientes originados desde el servidor.", "PRINCIPIO DE SEGURIDAD PERIMETRAL")

    # -------------------------------------------------------------
    # 2. CONEXIÓN SSH Y ADMINISTRACIÓN DESDE LAPTOPS MUNICIPALES
    # -------------------------------------------------------------
    add_h1("2. Conexión SSH y Gestión Administrativa desde Laptops Municipales")
    add_p("La administración del servidor y el monitoreo de Cisternin no requieren encontrarse físicamente en el rack municipal ni utilizar conexiones remotas inseguras. A través de la red privada, las laptops institucionales autorizadas operan como si estuvieran conectadas por cable directo al servidor.")

    add_h2("2.1 Preparación de la Laptop Municipal (Jefatura / Administrador)")
    add_bullet("Descargar e instalar el cliente oficial de Tailscale compatible con el sistema operativo de la laptop (Windows, macOS o Linux).")
    add_bullet("Iniciar sesión en el cliente utilizando la cuenta institucional autorizada por el área de informática (Google Workspace o Microsoft Entra ID municipal).")
    add_bullet("Verificar que la aplicación indique estado conectado. A partir de este momento, la laptop ingresa a la red privada segura de la municipalidad.")

    add_h2("2.2 Conexión Remota por Terminal SSH")
    add_p("Cada servidor registrado en la red privada recibe una dirección IP virtual fija dentro del rango 100.x.y.z. Esta dirección es accesible únicamente por los equipos autorizados dentro de la misma cuenta.")
    add_p("Para iniciar una sesión de terminal administrativa desde Windows Terminal, PowerShell o consola Linux/macOS, ejecute:")
    add_code_block(doc, "ssh usuario_servidor@100.x.y.z\n# Ejemplo:\nssh admin@100.85.12.34")
    add_p("Una vez autenticado mediante contraseña o llave SSH, dispondrá de acceso a la consola de comandos del sistema para gestionar los servicios de producción.")

    add_h2("2.3 Acceso a las Consolas Administrativas Web")
    add_p("Los módulos internos del sistema se encuentran restringidos a la red privada. Desde el navegador de la laptop municipal, se accede directamente a:")
    add_bullet(" Monitoreo de uso de memoria RAM, carga de CPU, inspección de almacenamiento de base vectorial y visualización de conversaciones.", "Panel de Servidor: http://100.x.y.z:8000/admin —")
    add_bullet(" Métricas de preguntas frecuentes ciudadanas, inventario de documentos oficiales indexados, descarga de reportes en CSV/JSON y módulo de actualización de contraseñas.", "Portal de Documentación y Reportes: http://100.x.y.z:8000/docs —")

    # -------------------------------------------------------------
    # 3. PROCEDIMIENTO DE DESPLIEGUE EN EL SERVIDOR
    # -------------------------------------------------------------
    add_h1("3. Procedimientos de Despliegue e Instalación en el Servidor")
    add_p("Se presentan dos métodos de instalación: la Opción 1 utiliza el script de automatización desarrollado para el proyecto, mientras que la Opción 2 detalla los comandos manuales paso a paso para entornos corporativos con políticas restrictivas de ejecución de scripts.")

    add_h2("3.1 Opción 1: Despliegue Automatizado mediante Script (Recomendada)")
    add_p("El proyecto provee el script de despliegue automatizado scripts/setup_red_cisternin.sh, el cual verifica dependencias, instala los paquetes oficiales, inicializa los servicios del sistema y solicita el token de conexión.")
    add_p("Procedimiento de ejecución en el servidor municipal:")
    add_code_block(doc, "cd /ruta/del/proyecto/Chatbot-RAG\nsudo bash scripts/setup_red_cisternin.sh")
    add_p("Comportamiento del script automatizado:")
    add_bullet("Comprueba la presencia de Tailscale; si no existe, descarga e instala el repositorio oficial e inicia el servicio de sistema tailscaled.")
    add_bullet("Si el servidor no está enlazado a la red, genera el enlace de autenticación para vincular la máquina.")
    add_bullet("Detecta la arquitectura del procesador (x86_64 o arm64) e instala el binario oficial de cloudflared.")
    add_bullet("Solicita el Token del túnel de Cloudflare e instala el servicio persistente del sistema, asociando el tráfico al puerto 8000.")
    add_bullet("Imprime un informe en pantalla con la IP privada asignada y las URLs de gestión.")

    add_h2("3.2 Opción 2: Despliegue Manual Paso a Paso (Alternativa sin Script)")
    add_p("En caso de que las políticas internas de la municipalidad prohíban la ejecución de instaladores automatizados, siga estrictamente los siguientes pasos manuales:")

    add_p("Paso A: Instalación manual de Tailscale (VPN Privada)", "Subpaso 1: ")
    add_code_block(doc, "# 1. Añadir la llave de firma oficial del repositorio de Tailscale (Debian/Ubuntu)\ncurl -fsSL https://pkgs.tailscale.com/stable/ubuntu/jammy.noarmor.gpg | sudo tee /usr/share/keyrings/tailscale-archive-keyring.gpg >/dev/null\ncurl -fsSL https://pkgs.tailscale.com/stable/ubuntu/jammy.tailscale-keyring.list | sudo tee /etc/apt/sources.list.d/tailscale.list\n\n# 2. Actualizar lista de paquetes e instalar\nsudo apt-get update\nsudo apt-get install -y tailscale\n\n# 3. Habilitar e iniciar el servicio persistente en systemd\nsudo systemctl enable --now tailscaled")

    add_p("Paso B: Autenticación y vinculación del servidor a la red privada", "Subpaso 2: ")
    add_code_block(doc, "# Ejecutar comando de registro asignando el nombre de host institucional\nsudo tailscale up --hostname=\"cisternin-server\" --accept-routes\n\n# Obtener y anotar la dirección IP asignada dentro de la red privada\ntailscale ip -4")

    add_p("Paso C: Instalación manual de Cloudflare Tunnel (cloudflared)", "Subpaso 3: ")
    add_code_block(doc, "# 1. Descargar el paquete oficial (.deb para arquitectura x86_64 / amd64)\ncurl -L --output /tmp/cloudflared.deb https://github.com/cloudflare/cloudflared/releases/latest/download/cloudflared-linux-amd64.deb\n\n# 2. Instalar el paquete en el sistema\nsudo dpkg -i /tmp/cloudflared.deb\nrm /tmp/cloudflared.deb\n\n# 3. Verificar instalación correcta de la herramienta\ncloudflared --version")

    add_p("Paso D: Vinculación del túnel de producción como servicio del sistema", "Subpaso 4: ")
    add_code_block(doc, "# Reemplazar <CLOUDFLARED_TOKEN> por el token generado en el panel Cloudflare Zero Trust\nsudo cloudflared service install <CLOUDFLARED_TOKEN>\n\n# Habilitar e iniciar el servicio en systemd\nsudo systemctl daemon-reload\nsudo systemctl enable --now cloudflared\n\n# Comprobar estado del servicio\nsudo systemctl status cloudflared")

    add_p("Paso E: Verificación de disponibilidad del servicio Cisternin", "Subpaso 5: ")
    add_code_block(doc, "# Comprobar que el backend de Cisternin responde en el puerto local 8000\ncurl -I http://127.0.0.1:8000/\n\n# Si se ejecuta mediante Docker Compose:\ndocker compose up -d\n\n# Si se ejecuta mediante servicio nativo:\nsudo systemctl restart cisternin")

    # -------------------------------------------------------------
    # 4. REQUERIMIENTOS PARA EL EQUIPO WEB
    # -------------------------------------------------------------
    add_h1("4. Requerimientos para el Área de Desarrollo Web Municipal")
    add_p("Una de las fortalezas centrales de este diseño es que opera como un microservicio completamente desacoplado. El equipo a cargo del portal municipal (React, Strapi o HTML tradicional) no requiere realizar modificaciones a su infraestructura de base de datos ni desplegar entornos de Python.")
    add_p("Se solicitan únicamente dos acciones puntuales de integración:")

    add_h2("4.1 Requerimiento A: Registro DNS Institucional (1 minuto)")
    add_p("En el panel de administración de DNS municipal (o zona de Cloudflare institucional), crear un registro de tipo CNAME que apunte el subdominio deseado hacia el identificador del túnel:")
    
    table_dns = doc.add_table(rows=5, cols=2)
    table_dns.alignment = WD_TABLE_ALIGNMENT.CENTER
    dns_rows = [
        ("Tipo de Registro:", "CNAME"),
        ("Nombre / Host:", "chatbot (quedando como chatbot.lacisterna.cl)"),
        ("Valor / Destino:", "<ID_DEL_TUNEL>.cfargotunnel.com"),
        ("Estado de Proxy:", "Proxied / Nube Naranja activada (Cifrado TLS automático)"),
        ("TTL:", "Automático")
    ]
    for idx, (label, val) in enumerate(dns_rows):
        cl = table_dns.rows[idx].cells[0]
        cr = table_dns.rows[idx].cells[1]
        cl.width = Inches(2.2)
        cr.width = Inches(4.3)
        set_cell_background(cl, "F8FAFC")
        set_cell_background(cr, "FFFFFF")
        set_cell_margins(cl, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cr, top=60, bottom=60, left=100, right=100)
        
        pl = cl.paragraphs[0]
        rl = pl.add_run(label)
        rl.bold = True
        rl.font.name = "Arial"
        rl.font.size = Pt(9.5)
        
        pr = cr.paragraphs[0]
        rr = pr.add_run(val)
        rr.font.name = "Arial"
        rr.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("4.2 Requerimiento B: Integración del Widget en el Portal Web (2 minutos)")
    add_p("Para habilitar la ventana de diálogo interactiva de Cisternin en la esquina inferior derecha del portal municipal, el equipo de desarrollo web debe incluir la siguiente etiqueta de script en el pie de página (footer) de su plantilla global:")
    add_code_block(doc, "<!-- Widget Oficial Cisternin - Municipalidad de La Cisterna -->\n<script \n  src=\"https://chatbot.lacisterna.cl/static/widget.js\" \n  data-bot-url=\"https://chatbot.lacisterna.cl/chat\" \n  defer>\n</script>")

    add_h2("4.3 Características de Aislamiento Técnico del Widget")
    add_bullet("El script tiene un peso inferior a 6 KB y se carga de forma asíncrona diferida (defer), garantizando cero impacto en el puntaje de rendimiento de la página web (Core Web Vitals).", "Cero impacto en rendimiento: ")
    add_bullet("El chat se renderiza en un elemento Iframe aislado con Sandbox perimetral. Los estilos CSS y scripts de Cisternin no pueden alterar ni colisionar con el diseño visual del sitio municipal.", "Aislamiento total de estilos (CSS Sandbox): ")
    add_bullet("El widget se adapta dinámicamente a pantallas de escritorio, tablets y dispositivos móviles con soporte táctil nativo.", "Diseño responsivo: ")

    # -------------------------------------------------------------
    # 5. MATRIZ DE CONFIGURACIÓN DE VARIABLES (.ENV)
    # -------------------------------------------------------------
    add_h1("5. Matriz de Variables de Configuración del Servidor")
    add_p("Todas las configuraciones operativas del sistema se centralizan en el archivo .env ubicado en la raíz del proyecto. Para editarlas en consola, utilice:")
    add_code_block(doc, "nano .env")
    add_p("A continuación se detallan las variables principales y sus propósitos:")

    table_env = doc.add_table(rows=6, cols=3)
    table_env.alignment = WD_TABLE_ALIGNMENT.CENTER
    env_headers = ["Variable de Entorno", "Valor por Defecto / Ejemplo", "Descripción Operativa"]
    for i, h in enumerate(env_headers):
        cell = table_env.rows[0].cells[i]
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    env_data = [
        ("API_KEY", "AlzaSyD-xxxxxxxx...", "Clave de acceso al proveedor de Inteligencia Artificial (Google Gemini o compatible)."),
        ("DOMAIN_NAME", "chatbot.lacisterna.cl", "Subdominio institucional asignado para la publicación web mediante túnel."),
        ("LLM_PROVIDER", "google (o ollama)", "Selecciona entre API en la nube (google) o ejecución 100% local en servidor con GPU (ollama)."),
        ("ADMIN_PANEL_CLAVE", "SuperSecreto...", "Contraseña maestra de acceso al panel de administración del servidor (/admin)."),
        ("DOCS_PASSWORD", "docs2026", "Contraseña de acceso al portal de documentación, métricas y reportes (/docs).")
    ]
    for idx, (var, ex, desc) in enumerate(env_data, start=1):
        row = table_env.rows[idx]
        for c_idx, val in enumerate([var, ex, desc]):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F8FAFC" if idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Consolas" if c_idx == 0 else "Arial"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) if c_idx == 0 else RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # -------------------------------------------------------------
    # 6. PROTOCOLO DE CONTINGENCIA Y DIAGNÓSTICO
    # -------------------------------------------------------------
    add_h1("6. Protocolo de Diagnóstico y Contingencia Operativa")
    add_p("A continuación se resumen las acciones de verificación recomendadas ante incidencias comunes de conectividad:")
    
    add_bullet("Comprobar el estado del proceso en el servidor ejecutando 'sudo systemctl status tailscaled'. Si el servicio está activo pero la laptop no conecta, verificar que ambos dispositivos se encuentren autenticados bajo la misma cuenta corporativa.", "Incidencia: La laptop no conecta al servidor vía Tailscale — ")
    add_bullet("Verificar los registros de conexión del túnel con 'sudo journalctl -u cloudflared -n 50 --no-pager'. Comprobar que el servicio local de Cisternin esté respondiendo en el puerto 8000 mediante 'curl -I http://127.0.0.1:8000/'.", "Incidencia: El subdominio chatbot.lacisterna.cl arroja error 502 / Bad Gateway — ")
    add_bullet("El sistema dispone de conmutación por error (fallback) automática: si el servidor local de IA no responde en 5 segundos, conmuta automáticamente a la API en la nube para mantener la atención ciudadana ininterrumpida.", "Incidencia: Caída momentánea del motor local de IA — ")

    # Firma
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.keep_with_next = True
    r_s1 = p_sign.add_run("ÁREA DE DESARROLLO E INFRAESTRUCTURA TI\n")
    r_s1.bold = True
    r_s1.font.name = "Arial"
    r_s1.font.size = Pt(10)
    r_s1.font.color.rgb = RGBColor(0x0F, 0x29, 0x4A)
    
    r_s2 = p_sign.add_run("Proyecto Asistente Virtual Ciudadano Cisternin\nIlustre Municipalidad de La Cisterna — Octubre de 2026")
    r_s2.font.name = "Arial"
    r_s2.font.size = Pt(9.5)
    r_s2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.save(filename)
    print(f"Documento Word guardado exitosamente: {filename}")

if __name__ == "__main__":
    out_docx = "/home/aspen/Alcaldia_Practica/docs/Guia_Tecnica_Red_y_Despliegue_Cisternin.docx"
    build_docx(out_docx)
    
    # Copiar también a static/docs/ para descarga directa
    static_docx = "/home/aspen/Alcaldia_Practica/static/docs/Guia_Tecnica_Red_y_Despliegue_Cisternin.docx"
    os.makedirs(os.path.dirname(static_docx), exist_ok=True)
    subprocess.run(["cp", out_docx, static_docx], check=True)
    
    # Convertir a PDF con LibreOffice
    print("Convirtiendo a PDF con LibreOffice...")
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to", "pdf",
        out_docx,
        "--outdir", "/home/aspen/Alcaldia_Practica/docs"
    ]
    subprocess.run(cmd, check=True)
    
    # Copiar PDF a static/docs/
    out_pdf = "/home/aspen/Alcaldia_Practica/docs/Guia_Tecnica_Red_y_Despliegue_Cisternin.pdf"
    static_pdf = "/home/aspen/Alcaldia_Practica/static/docs/Guia_Tecnica_Red_y_Despliegue_Cisternin.pdf"
    subprocess.run(["cp", out_pdf, static_pdf], check=True)
    print("PDF generado y copiado exitosamente a static/docs/")
