import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.line_spacing = 1.15

    # --- ENCABEZADO INSTITUCIONAL ---
    p_inst = doc.add_paragraph()
    r_inst = p_inst.add_run("ILUSTRE MUNICIPALIDAD DE LA CISTERNA\nDIRECCIÓN DE INFORMÁTICA Y TECNOLOGÍA\nPROYECTO ASISTENTE VIRTUAL CIUDADANO (CISTERNIN)")
    r_inst.font.size = Pt(9)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    # TÍTULO DEL DOCUMENTO
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("TÉRMINOS DE REQUERIMIENTO TÉCNICO Y FINANCIERO")
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("ADQUISICIÓN DE SERVICIO DE API DE INTELIGENCIA ARTIFICIAL GENERATIVA (LLM) PARA ATENCIÓN VECINAL 24/7")
    r_sub.font.size = Pt(11.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # METADATOS EN TABLA
    tbl_meta = doc.add_table(rows=4, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("DESTINATARIO:", "Administración Municipal / Dirección de Finanzas / Jefatura de Informática"),
        ("UNIDAD TÉCNICA:", "Equipo de Desarrollo e Innovación Tecnológica (Proyecto Cisternin)"),
        ("FECHA:", "Octubre de 2026"),
        ("OBJETO ÚNICO:", "Contratación de consumo de API de Modelo de Lenguaje (LLM) empresarial")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = tbl_meta.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.0), Inches(4.9)
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=70, bottom=70)
        set_cell_margins(c1, top=70, bottom=70)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    # 1. ANTECEDENTES Y ESTADO DEL SISTEMA
    h1 = doc.add_paragraph()
    r = h1.add_run("1. ANTECEDENTES Y ESTADO OPERATIVO DEL SISTEMA")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph()
    p.add_run("La Municipalidad de La Cisterna dispone de la solución integral de Asistente Virtual Ciudadano (\"Cisternin\"), desarrollada y lista para su despliegue en producción. ")
    p.add_run("Toda la arquitectura técnica, plataformas de administración y base de conocimiento comunal ya se encuentran completamente construidas e integradas, sin requerir inversión en desarrollo de software ni compras de equipamiento físico adicional.")

    p_check = doc.add_paragraph()
    p_check.paragraph_format.left_indent = Inches(0.2)
    p_check.add_run("• Componentes ya operativos al 100% ($0 CLP requeridos):\n").bold = True
    p_check.add_run("  - Motor conversacional RAG (FastAPI / Python) con búsqueda semántica local.\n")
    p_check.add_run("  - Base de datos vectorial municipal (ChromaDB) con 583 fragmentos oficiales indexados (trámites, licencias, farmacia, veterinaria, Tarjeta Vecino, aseo y autoridades).\n")
    p_check.add_run("  - Portal público de atención ciudadana y widget web flotante para inserción en www.cisterna.cl.\n")
    p_check.add_run("  - Panel de administración (/admin), portal de documentación con reportes en Excel (/docs) y canal de derivación oficial a WhatsApp Municipal.\n")
    p_check.add_run("  - Túnel seguro de alta disponibilidad perimetral (sin abrir puertos en el firewall municipal).")

    # 2. DEFINICIÓN DEL REQUERIMIENTO
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    r = h2.add_run("2. DEFINICIÓN DEL REQUERIMIENTO: SERVICIO DE INTELIGENCIA ARTIFICIAL")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph()
    p.add_run("El único componente que se solicita adquirir es el ")
    p.add_run("Servicio de API de Modelo de Lenguaje (LLM de grado empresarial)").bold = True
    p.add_run(", necesario para que el bot genere las respuestas conversacionales a los vecinos de forma ininterrumpida y con tiempo de respuesta inferior a 1,5 segundos.")

    p_req = doc.add_paragraph()
    p_req.paragraph_format.left_indent = Inches(0.2)
    p_req.add_run("• Especificaciones técnicas requeridas (Neutralidad Tecnológica - Ley 19.886):\n").bold = True
    p_req.add_run("  1. Protocolo de integración: API estándar RESTful HTTPS (JSON entrada/salida).\n")
    p_req.add_run("  2. Capacidad y latencia: Modelo optimizado de alta velocidad (Time to First Token < 1,5s).\n")
    p_req.add_run("  3. Ventana de contexto: Mínimo 128.000 tokens para análisis integral de decretos y ordenanzas.\n")
    p_req.add_run("  4. Privacidad y Seguridad (Ley N° 19.628): El proveedor no debe retener ni utilizar las consultas de los vecinos para entrenamiento de modelos públicos.\n")
    p_req.add_run("  5. Modelo de cobro: Pago por uso exacto (Pay-As-You-Go) con control de gasto máximo mensual (Budget Cap).")

    # 3. MODELO DE COSTOS Y PRESUPUESTO
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    r = h3.add_run("3. MODELO DE COSTOS Y PRESUPUESTO SOLICITADO")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph()
    p.add_run("El costo real de cada interacción vecinal es sumamente bajo debido a la eficiencia del motor RAG municipal, que inyecta únicamente los fragmentos necesarios. ")
    p.add_run("El costo directo estimado es de apenas ")
    p.add_run("$0,14 CLP por consulta ciudadana.").bold = True

    # TABLA COMPARATIVA
    tbl_costo = doc.add_table(rows=5, cols=4)
    tbl_costo.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Concepto Financiero", "Monto Mensual (CLP)", "Monto Anual (CLP)", "Proyección 2 Años (24m)"]
    for j, h in enumerate(headers):
        cell = tbl_costo.rows[0].cells[j]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=100, bottom=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    filas = [
        ("Consumo Base Estimado (1.000 consultas/día)", "$4.275 CLP (~$4,50 USD)", "$51.300 CLP", "$102.600 CLP"),
        ("Presupuesto Solicitado (Techo Operativo)", "$100.000 CLP", "$1.200.000 CLP (1,2 M)", "$2.400.000 CLP (2,4 M)"),
        ("Capacidad de Atención Cubierta", "Hasta 701.750 consultas/mes", "Hasta 8,4 Millones/año", "Hasta 16,8 Millones"),
        ("Costo por Consulta Vecinal", "$0,14 CLP / consulta", "$0,14 CLP / consulta", "$0,14 CLP / consulta")
    ]
    for i, data in enumerate(filas, start=1):
        row = tbl_costo.rows[i]
        bg = "F9FBFD" if i % 2 == 1 else "FFFFFF"
        if i == 2:
            bg = "EBF5FB"  # Destacar presupuesto solicitado
        for j, val in enumerate(data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90)
            p = cell.paragraphs[0]
            if j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if i == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph().paragraph_format.space_before = Pt(6)

    # 4. JUSTIFICACIÓN DEL PRESUPUESTO
    h4 = doc.add_paragraph()
    r = h4.add_run("4. JUSTIFICACIÓN DE LA ASIGNACIÓN DE $100.000 CLP/MES ($1,2 MILLONES/AÑO)")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p = doc.add_paragraph()
    p.add_run("Si el consumo base promedio para 1.000 atenciones diarias representa ~$4.275 CLP mensuales, se justifica solicitar un presupuesto de ")
    p.add_run("$100.000 CLP mensuales ($1.200.000 CLP anuales)").bold = True
    p.add_run(" por los siguientes motivos operativos y de gestión pública:")

    p_just = doc.add_paragraph()
    p_just.paragraph_format.left_indent = Inches(0.2)
    p_just.add_run("1. Variabilidad por vecino: ").bold = True
    p_just.add_run("Un vecino no siempre realiza una sola pregunta; comúnmente formula consultas sucesivas (3 a 8 preguntas) para resolver dudas de subsidios, requisitos vehiculares o licencias de conducir.\n")
    p_just.add_run("2. Picos estacionales de alta demanda: ").bold = True
    p_just.add_run("Durante períodos de alta afluencia (marzo por Permisos de Circulación, mayo por patentes comerciales), el tráfico municipal se multiplica hasta por cinco. El presupuesto asegura capacidad para más de 700.000 consultas mensuales sin interrupción.\n")
    p_just.add_run("3. Garantía de Cero Caídas (Zero-Downtime): ").bold = True
    p_just.add_run("Un crédito de $100.000 CLP/mes garantiza que el servicio nunca se suspenda por 'cuota excedida' o límites de saldo.\n")
    p_just.add_run("4. Eficiencia y protección del erario municipal: ").bold = True
    p_just.add_run("Bajo el modelo Pay-As-You-Go, el municipio solo paga por las consultas efectivamente realizadas. Si en un mes se consumen $12.000 CLP, los $88.000 CLP restantes no se pierden ni se cobran, quedando resguardados por el límite mensual configurado.")

    # 5. PROCEDIMIENTO DE HABILITACIÓN
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(8)
    r = h5.add_run("5. PROCEDIMIENTO DIRECTO PARA LA PUESTA EN MARCHA")
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_proc = doc.add_paragraph()
    p_proc.paragraph_format.left_indent = Inches(0.2)
    p_proc.add_run("Paso 1: ").bold = True
    p_proc.add_run("Aprobación del requerimiento presupuestario de $100.000 CLP/mes ($1.200.000 CLP/año) para servicio de API de IA.\n")
    p_proc.add_run("Paso 2: ").bold = True
    p_proc.add_run("Habilitación de la cuenta de nube municipal con límite de gasto (Budget Cap) configurado en $100.000 CLP/mes.\n")
    p_proc.add_run("Paso 3: ").bold = True
    p_proc.add_run("Entrega de la credencial API al equipo técnico para su vinculación con el bot (proceso de 5 minutos).\n")
    p_proc.add_run("Paso 4: ").bold = True
    p_proc.add_run("Inserción del widget en www.cisterna.cl y habilitación del servicio ciudadano 24/7.")

    # FIRMAS
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(28)
    p_sign.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sign = p_sign.add_run("_________________________________________\nMATHIAS ARRIAGADA\nDesarrollo e Innovación — Proyecto Cisternin\nIlustre Municipalidad de La Cisterna")
    r_sign.font.size = Pt(10)
    r_sign.font.bold = True

    # Guardar DOCX
    docx_path = "/home/aspen/Requerimiento_Adquisicion_IA_Cisternin.docx"
    doc.save(docx_path)
    print(f"Documento DOCX generado: {docx_path}")

    # Compilar a PDF con LibreOffice
    cmd = ["libreoffice", "--headless", "--convert-to", "pdf", docx_path, "--outdir", "/home/aspen"]
    subprocess.run(cmd, check=True)
    pdf_path = "/home/aspen/Requerimiento_Adquisicion_IA_Cisternin.pdf"
    print(f"Documento PDF generado: {pdf_path}")

    # Copiar a docs/ y static/docs/
    os.makedirs("/home/aspen/Alcaldia_Practica/docs", exist_ok=True)
    os.makedirs("/home/aspen/Alcaldia_Practica/static/docs", exist_ok=True)
    subprocess.run(["cp", docx_path, "/home/aspen/Alcaldia_Practica/docs/"], check=True)
    subprocess.run(["cp", pdf_path, "/home/aspen/Alcaldia_Practica/docs/"], check=True)
    subprocess.run(["cp", pdf_path, "/home/aspen/Alcaldia_Practica/static/docs/"], check=True)
    subprocess.run(["cp", docx_path, "/home/aspen/Alcaldia_Practica/static/docs/"], check=True)
    print("Archivos copiados a docs/ y static/docs/")

if __name__ == "__main__":
    create_document()
